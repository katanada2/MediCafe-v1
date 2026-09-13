import csv
import io
import zipfile
from dataclasses import dataclass
from datetime import date

from docx import Document

from .domain import CommandError


PARSER_VERSION = "f1-v1"
EXPECTED_HEADER = ["row_id", "patient_ref", "service_date", "note"]
MAX_ROWS = 1_000
MAX_DOCX_EXPANDED = 25 * 1024 * 1024


@dataclass(frozen=True)
class ParsedRow:
    row_ordinal: int
    raw_values: dict
    normalized_values: dict
    warnings: list
    source_locator: dict


def _normalized(raw):
    values = {key: value.strip() for key, value in raw.items()}
    warnings = []
    if not values["patient_ref"]:
        warnings.append("patient_ref_blank")
    if not values["service_date"]:
        warnings.append("service_date_blank")
    else:
        try:
            date.fromisoformat(values["service_date"])
        except ValueError:
            warnings.append("service_date_invalid")
    return values, warnings


def _rows_with_duplicate_warnings(rows):
    counts = {}
    for row in rows:
        row_id = row.normalized_values["row_id"]
        counts[row_id] = counts.get(row_id, 0) + 1
    return [ParsedRow(
        row.row_ordinal,
        row.raw_values,
        row.normalized_values,
        row.warnings + (["duplicate_row_id"] if counts[row.normalized_values["row_id"]] > 1 else []),
        row.source_locator,
    ) for row in rows]


def parse_csv(content):
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise CommandError("csv_not_utf8") from exc
    reader = csv.DictReader(io.StringIO(text, newline=""))
    if reader.fieldnames != EXPECTED_HEADER:
        raise CommandError("invalid_header")
    rows = []
    for ordinal, record in enumerate(reader, start=1):
        if ordinal > MAX_ROWS:
            raise CommandError("row_limit_exceeded")
        if None in record:
            raise CommandError("invalid_column_count")
        raw = {key: record[key] for key in EXPECTED_HEADER}
        normalized, warnings = _normalized(raw)
        rows.append(ParsedRow(ordinal, raw, normalized, warnings, {"format": "csv", "data_row": ordinal}))
    return _rows_with_duplicate_warnings(rows)


def parse_docx(content):
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if any(item.flag_bits & 0x1 for item in archive.infolist()):
                raise CommandError("docx_encrypted")
            if sum(item.file_size for item in archive.infolist()) > MAX_DOCX_EXPANDED:
                raise CommandError("docx_expanded_limit_exceeded")
    except zipfile.BadZipFile as exc:
        raise CommandError("docx_malformed") from exc
    try:
        document = Document(io.BytesIO(content))
    except Exception as exc:
        raise CommandError("docx_malformed") from exc
    if len(document.tables) != 1:
        raise CommandError("docx_table_shape_unsupported")
    table = document.tables[0]
    if not table.rows or [cell.text.strip() for cell in table.rows[0].cells] != EXPECTED_HEADER:
        raise CommandError("invalid_header")
    rows = []
    for ordinal, table_row in enumerate(table.rows[1:], start=1):
        if ordinal > MAX_ROWS:
            raise CommandError("row_limit_exceeded")
        if len(table_row.cells) != len(EXPECTED_HEADER):
            raise CommandError("invalid_column_count")
        raw = {key: cell.text for key, cell in zip(EXPECTED_HEADER, table_row.cells, strict=True)}
        normalized, warnings = _normalized(raw)
        rows.append(ParsedRow(
            ordinal, raw, normalized, warnings,
            {"format": "docx", "table": 1, "table_row": ordinal + 1},
        ))
    return _rows_with_duplicate_warnings(rows)


def parse(content, media_type):
    if media_type == "text/csv":
        return parse_csv(content)
    if media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return parse_docx(content)
    raise CommandError("media_type_unsupported")
