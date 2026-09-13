"""Small, deliberately synthetic F1 input builders.

These helpers keep fixtures in memory so the acceptance suite does not add
tracked artifacts or rely on a particular filesystem layout.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Iterable, Mapping

from docx import Document


CSV_MEDIA_TYPE = "text/csv"
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
CSV_HEADER = ("row_id", "patient_ref", "service_date", "note")
SYNTHETIC_SENTINEL = "SYNTHETIC_SENTINEL_F1_7F1"


def synthetic_rows(*, patient_ref: str = "000042", service_date: str = "2026-01-15",
                   note: str = SYNTHETIC_SENTINEL, row_id: str = "row-001") -> list[dict[str, str]]:
    return [{
        "row_id": row_id,
        "patient_ref": patient_ref,
        "service_date": service_date,
        "note": note,
    }]


def csv_bytes(rows: Iterable[Mapping[str, object]] | None = None, *, bom: bool = False,
              header: Iterable[str] = CSV_HEADER) -> bytes:
    rows = list(rows if rows is not None else synthetic_rows())
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=list(header), lineterminator="\n")
    writer.writeheader()
    for row in rows:
        writer.writerow({key: row.get(key, "") for key in header})
    encoded = stream.getvalue().encode("utf-8")
    return (b"\xef\xbb\xbf" if bom else b"") + encoded


def docx_bytes(rows: Iterable[Mapping[str, object]] | None = None, *, table_count: int = 1,
               header: Iterable[str] = CSV_HEADER) -> bytes:
    rows = list(rows if rows is not None else synthetic_rows())
    document = Document()
    for table_index in range(table_count):
        table = document.add_table(rows=1, cols=len(tuple(header)))
        for cell, value in zip(table.rows[0].cells, header, strict=True):
            cell.text = str(value)
        if table_index == 0:
            for row in rows:
                cells = table.add_row().cells
                for cell, key in zip(cells, header, strict=True):
                    cell.text = str(row.get(key, ""))
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def empty_docx_bytes() -> bytes:
    stream = io.BytesIO()
    Document().save(stream)
    return stream.getvalue()

